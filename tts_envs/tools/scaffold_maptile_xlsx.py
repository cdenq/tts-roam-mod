#!/usr/bin/env python3
"""Create or top up data/maptile.xlsx from what build_map_tiles.py produces.

One row per map image, so 1,200+ URL cells never have to be typed as keys by
hand. Safe to re-run at any time: rows are matched by `image` name, and every
cell you have already filled (url, description, weights, ...) is kept. New
biomes, patterns or densities in asset_variables.csv just add rows.

  sheet                 one row per                       key
  settlement .. scarred map image of that deck            biome > pattern > orientation > density
  backs                 deck back image                   deck
  biome_descriptions    biome                             biome (= tile GMNotes)
  maptile_table         biome: what Render spawns         gmnotes
  density               density stage                     density -> explorations
  weights               draw roll                         param -> pool ("a|b|b" = b twice as likely)

Rows the generator no longer describes are reported, never deleted.
Within each biome, deck-sheet rows are in FILE-NAME order (as a file browser
lists the biome's folder), so URLs copied in folder order paste straight in.

Run: python tools/scaffold_maptile_xlsx.py
"""
import os, sys

sys.dont_write_bytecode = True
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from datasets import MAPTILES, ASSET_VARS
import build_map_tiles as bm

try:
  import openpyxl
  from openpyxl.styles import Font, PatternFill
except ImportError:
  sys.exit("scaffold_maptile_xlsx.py needs openpyxl: pip install openpyxl")

DECK_COLUMNS = ["biome", "pattern", "orientation", "density", "image", "file", "url"]
SHEETS = {
  "backs": ["deck", "name", "image", "file", "url"],
  "biome_descriptions": ["biome", "deck", "name", "description"],
  # placeholder headers -- see plans/auto_TODO.md "Maptile Render table"
  "maptile_table": ["gmnotes", "spawnmeeple1", "spawnmeeple2", "spawn1", "spawn2", "spawngold", "notes"],
  "density": ["density", "explorations", "notes"],
  "weights": ["param", "pool", "notes"],
}
# Display names that .title() would get wrong.
BIOME_NAMES = {
  "blossomgrove": "Blossom Grove", "mushroomhollow": "Mushroom Hollow", "briarthicket": "Briar Thicket",
  "wildflowermeadow": "Wildflower Meadow", "waterfallcliffs": "Waterfall Cliffs",
  "standingstones": "Standing Stones", "hotsprings": "Hot Springs", "lumbercamp": "Lumber Camp",
}
HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="3B2B20")
WIDTHS = {"url": 60, "description": 60, "file": 42, "image": 30, "pool": 70, "notes": 50}


def sheet(wb, title, headers):
  """The sheet, created with a styled header if missing; headers checked if present."""
  if title in wb.sheetnames:
    ws = wb[title]
    have = [c.value for c in ws[1]]
    for h in headers:  # add any missing column at the end, never reorder
      if h not in have:
        ws.cell(1, len(have) + 1, h)
        have.append(h)
  else:
    ws = wb.create_sheet(title)
    for i, h in enumerate(headers, 1):
      ws.cell(1, i, h)
  for i, c in enumerate(ws[1], 1):
    c.font, c.fill = HEADER_FONT, HEADER_FILL
    ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = WIDTHS.get(c.value, 14)
  ws.freeze_panes = "A2"
  return ws


def upsert(ws, key, rows):
  """Add rows whose `key` value is not on the sheet yet; fill only EMPTY cells
  of existing rows. Returns (added, unknown keys already on the sheet)."""
  header = [c.value for c in ws[1]]
  col = {h: i for i, h in enumerate(header)}
  existing = {}
  for r in ws.iter_rows(min_row=2):
    v = r[col[key]].value
    if v not in (None, ""):
      existing[str(v)] = r
  added = 0
  for row in rows:
    k = str(row[key])
    if k in existing:
      cells = existing[k]
      for h, v in row.items():
        if h in col and cells[col[h]].value in (None, "") and v not in (None, ""):
          cells[col[h]].value = v
    else:
      ws.append([row.get(h, "") for h in header])
      added += 1
  wanted = {str(r[key]) for r in rows}
  return added, sorted(set(existing) - wanted)


def main():
  glob, groups = bm.load_config(ASSET_VARS)
  bm.setup(bm.parse_global(glob))
  decks = bm.parse_decks({g: v for g, v in groups.items() if g.startswith("deck_")})
  regions = [bm.parse_region(n, raw) for n, raw in groups.items() if not n.startswith("deck_")]
  plan = bm.plan(regions, decks)

  wb = openpyxl.load_workbook(MAPTILES) if os.path.exists(MAPTILES) else openpyxl.Workbook()
  if "Sheet" in wb.sheetnames and len(wb.sheetnames) == 1:
    del wb["Sheet"]

  report = []
  deck_of = {r["name"]: r["deck"] for r in regions}
  for did in decks:
    ws = sheet(wb, did, DECK_COLUMNS)
    rows = [{"biome": t["region"], "pattern": t["pattern"], "orientation": t["edges"].replace("|", ""),
             "density": t["density"], "image": t["name"], "file": t["file"], "url": ""}
            for t in plan if t["pattern"] != "back" and deck_of[t["region"]] == did]
    # Folder order: within a biome, rows sort by image name -- the order a file
    # browser lists assets/maptiles/<biome>/, so a column of URLs copied from
    # the uploaded folder pastes straight in.
    biome_order = list(dict.fromkeys(r["biome"] for r in rows))
    rows.sort(key=lambda r: (biome_order.index(r["biome"]), r["image"]))
    report.append((did, ws, "image", rows))
  report.append(("backs", sheet(wb, "backs", SHEETS["backs"]), "deck",
                 [{"deck": d["id"], "name": d["name"], "image": "%s_back" % d["id"],
                   "file": "backs/%s_back.png" % d["id"], "url": ""} for d in decks.values()]))
  report.append(("biome_descriptions", sheet(wb, "biome_descriptions", SHEETS["biome_descriptions"]), "biome",
                 [{"biome": r["name"], "deck": r["deck"], "name": BIOME_NAMES.get(r["name"], r["name"].title()),
                   "description": ""} for r in regions]))
  report.append(("maptile_table", sheet(wb, "maptile_table", SHEETS["maptile_table"]), "gmnotes",
                 [{"gmnotes": r["name"]} for r in regions]))
  densities = sorted({d for r in regions for d in r["densities"]})
  report.append(("density", sheet(wb, "density", SHEETS["density"]), "density",
                 [{"density": d, "explorations": d, "notes": "placeholder: explorations = density"}
                  for d in densities]))
  patterns = list(dict.fromkeys(p for r in regions for p in r["patterns"]))
  pools = [{"param": "deck", "pool": "|".join(decks), "notes": "repeat an entry to make it likelier"}]
  pools += [{"param": "biome:%s" % d, "pool": "|".join(r["name"] for r in regions if r["deck"] == d)}
            for d in decks]
  pools.append({"param": "pattern", "pool": "|".join(patterns)})
  pools += [{"param": "orientation:%s" % p, "pool": "|".join("".join(o) for o in bm.orientations(p))}
            for p in patterns]
  pools.append({"param": "density", "pool": "|".join(str(d) for d in densities),
                "notes": "e.g. 1|1|2|2|3 makes density 3 half as likely"})
  report.append(("weights", sheet(wb, "weights", SHEETS["weights"]), "param", pools))

  order = list(decks) + list(SHEETS)
  wb._sheets.sort(key=lambda ws: order.index(ws.title) if ws.title in order else len(order))
  for title, ws, key, rows in report:
    added, stale = upsert(ws, key, rows)
    msg = "  %-20s %4d rows" % (title, ws.max_row - 1)
    if added:
      msg += ", %d added" % added
    if stale:
      msg += ", %d no longer generated (kept; delete by hand): %s" % (len(stale), ", ".join(stale[:5]))
    print(msg)
  wb.save(MAPTILES)
  print("wrote %s" % os.path.relpath(MAPTILES, os.path.dirname(MAPTILES)))


if __name__ == "__main__":
  main()
