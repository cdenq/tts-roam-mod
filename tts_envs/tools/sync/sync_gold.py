#!/usr/bin/env python3
"""Compute every gold value from data/gold_variables.csv and push it to source.

gold_variables.csv is the only place a gold number is authored by hand. Everything else
derives from the anchor, so rescaling the whole economy is one edit there.

This tool runs BACKWARDS compared to the other syncs: they read data/ and write
generated/, while this one writes INTO data/ --

  tiles.xlsx  items/skills/statuses  base_gold   computed; light-yellow
                                                header, like stat
  tiles.xlsx  lootbox_table        spawngold   a "lo-hi" range around the
                                                tier midpoint
  origins.xlsx                       spawngold   from the gold_tier column
  generated/economy_data.lua   ECONOMY -- margins + rep bands, for Global

A hand-typed value survives: an item whose base_gold does not match what the
rules would produce is reported as an OVERRIDE and left alone. That is why
base_gold is written in place rather than hidden in a lookup sheet.

Run: python tools/sync/sync_gold.py
     python tools/sync/sync_gold.py --check    report, write nothing
     python tools/sync/sync_gold.py --force    overwrite overrides too
"""
import csv, math, os, re, sys

sys.dont_write_bytecode = True
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(1, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))  # tools/: datasets, spawn_cells, lua_inline
from datasets import (ROOT, GOLD, TILES, ORIGINS, ITEM_SHEETS, LOOTBOX_SHEET,
                      ORIGINS_SHEET, load, read_rows, write_back, save)

OUT = os.path.join(ROOT, "generated", "economy_data.lua")
WORKSPACE = os.path.dirname(ROOT)
GLOBAL = os.path.join(WORKSPACE, ".tts", "objects", "Global.lua")

BEGIN = "--[[ BEGIN GENERATED ECONOMY DATA -- do not edit ]]"
END = "--[[ END GENERATED ECONOMY DATA ]]"

# The items sheet already marks computed columns light yellow (stat,
# damage_range) and authored ones grey. base_gold joins the yellow set.
COMPUTED_FILL = "FFFFF2CC"

REP_ORDER = ["enemy", "cautious", "neutral", "friendly", "revered"]

#[[ Every 0-3 tier ladder. ladder_for() picks one per row; validate() requires
# all four rungs of each, because a missing rung prices its items at nothing
# instead of failing. quest_REWARD_tier is the gold a quest PAYS OUT, which is
# a different thing from quest_ITEM_tier, the worth of a plot token.
TIER_LADDERS = ("goods_tier", "equip_tier", "consumable_tier", "skill_tier",
                "misc_tier", "quest_item_tier", "quest_reward_tier",
                "lootbox_tier")

# Gems and pearls are subsubtype=gem. They feed no crafter (see taxonomy.md
# "Loot is not gathered"), so resale is the whole point of finding one.
TREASURE_SUBSUBTYPES = {"gem"}

#[[ gmnotes is the machine-readable face of a row: an in-game script reads it
# instead of parsing the player-facing description. Two families are stamped.
#
#   goods/material   <subsubtype><tier>            ore2
#   weapon           dmg:<lo>-<hi>|<type>|<category>
#                                                  dmg:3-9|Physical|projectile
#
# Pipe-separated so a Lua split is one gmatch, and the keys are stable even
# when a field is empty. These MUST agree with sync_items_data.py, which owns
# the same rules for the description text -- see DAMAGE_TYPE_BY_CATEGORY and
# damage_range() there.
MISC_CATEGORY = "miscellaneous"
MISC_DAMAGE = (1, 4)
TIER0_LO, TIER0_HI = 1, 6
TIER_LO_STEP, TIER_HI_STEP = 2, 3
DAMAGE_TYPE_BY_CATEGORY = {
  "channeling": "Magical",
  "infusing": "Magical",
  "blocking": "Anti-Physical",
  "warding": "Anti-Magical",
}
DEFAULT_DAMAGE_TYPE = "Physical"


def damage_range(tier, hands, category):
  """The (lo, hi) a weapon rolls. Mirrors sync_items_data.damage_range."""
  if category == MISC_CATEGORY:
    return MISC_DAMAGE
  if tier is None:
    return None
  lo = TIER0_LO + TIER_LO_STEP * tier
  hi = TIER0_HI + TIER_HI_STEP * tier
  if category == "slashing":
    hi -= 1
  if hands == 2:
    lo, hi = 2 * lo, 2 * hi
  return (lo, hi)


def weapon_gmnotes(subtype, tier, hands):
  """The stat line an in-game script parses off a weapon."""
  dmg = damage_range(tier, hands, subtype)
  if not dmg:
    return None
  dtype = DAMAGE_TYPE_BY_CATEGORY.get(subtype, DEFAULT_DAMAGE_TYPE)
  return "dmg:%d-%d|%s|%s" % (dmg[0], dmg[1], dtype, subtype)

warnings_out = []
overrides_out = []
overwritten_out = []


def warn(msg):
  warnings_out.append(msg)


def read_csv(path, key):
  """Spreadsheets often save smart quotes as cp1252; try UTF-8 then fall back."""
  for enc in ("utf-8-sig", "cp1252", "latin-1"):
    try:
      with open(path, newline="", encoding=enc) as f:
        return [r for r in csv.DictReader(f) if (r.get(key) or "").strip()]
    except UnicodeDecodeError:
      continue
  raise SystemExit("cannot decode " + path)


def load_sheet(path):
  # Formulas, not cached values: this workbook is written back.
  return load(path, data_only=False)


def header_index(ws):
  return {h: i for i, h in enumerate(c.value for c in ws[1]) if h}


def save_sheet(wb, path):
  """Excel holds an exclusive lock on an open workbook, and openpyxl's
  PermissionError says nothing useful about why."""
  try:
    wb.save(path)
  except PermissionError:
    sys.exit("cannot write %s -- it is open in Excel. Close it and re-run."
             % os.path.relpath(path, ROOT))


def mark_computed(ws, col):
  from openpyxl.styles import PatternFill
  ws.cell(row=1, column=col).fill = PatternFill(
    "solid", start_color=COMPUTED_FILL, end_color=COMPUTED_FILL)


# ---- the rules ------------------------------------------------------------

def read_rules():
  """gold_variables.csv is long-format: group + key is the composite key."""
  rules = {}
  for r in read_csv(GOLD, "group"):
    group = (r.get("group") or "").strip()
    key = (r.get("key") or "").strip()
    raw = (r.get("value") or "").strip()
    if not group or not key:
      continue
    try:
      value = float(raw)
    except ValueError:
      raise SystemExit("gold_variables.csv: %s/%s has non-numeric value %r"
                       % (group, key, raw))
    rules.setdefault(group, {})[key] = value
  return rules


def validate(rules):
  """Invariants worth failing the build over."""
  for g in (("anchor", "equip_type", "equip_hands", "margin", "rep",
             "rep_buy", "rep_sell", "origin_tier", "lootbox", "treasure",
             "junk", "craft") + TIER_LADDERS):
    if g not in rules:
      raise SystemExit("gold_variables.csv: missing group %r" % g)

  anchor = rules["anchor"].get("gold_anchor")
  if not anchor or anchor <= 0:
    raise SystemExit("gold_variables.csv: gold_anchor must be positive")

  sell, buy = rules["margin"]["sell_pct"], rules["margin"]["buy_pct"]
  if sell <= -1:
    raise SystemExit("gold_variables.csv: sell_pct (%g) must be above -1 -- "
                     "at -1 a vendor pays nothing for anything" % sell)
  if sell >= buy:
    raise SystemExit("gold_variables.csv: sell_pct (%g) must be below buy_pct (%g) -- "
                     "otherwise buying and reselling is free money"
                     % (sell, buy))

  # Every rep band needs all three columns, and the minimums must ascend or
  # band lookup silently picks the wrong one.
  for name in REP_ORDER:
    for g in ("rep", "rep_buy", "rep_sell"):
      if name not in rules[g]:
        raise SystemExit("gold_variables.csv: %s is missing band %r" % (g, name))
  mins = [rules["rep"][n] for n in REP_ORDER]
  if mins != sorted(mins):
    raise SystemExit("gold_variables.csv: rep band minimums are not ascending: %s" % mins)

  # Selling a sealed box must return LESS than opening it, or nobody opens
  # one. lootbox/sell_pct is that ratio directly, so the check is just < 1.
  box_sell = rules["lootbox"].get("sell_pct")
  if box_sell is None:
    raise SystemExit("gold_variables.csv: lootbox/sell_pct is missing")
  if box_sell >= 1.0:
    raise SystemExit("gold_variables.csv: lootbox sell_pct (%g) must be below 1.0, or "
                     "selling a box beats opening it" % box_sell)
  if "gem_mult" not in rules.get("treasure", {}):
    raise SystemExit("gold_variables.csv: treasure/gem_mult is missing")

  # A missing rung prices its items at nothing instead of failing.
  for g in TIER_LADDERS:
    for t in range(4):
      if str(t) not in rules[g]:
        raise SystemExit("gold_variables.csv: %s has no tier %d" % (g, t))
  for t in range(1, 6):
    if str(t) not in rules["origin_tier"]:
      raise SystemExit("gold_variables.csv: origin_tier has no tier %d" % t)


def trade_multiplier(*offsets):
  """Margin, rep and shop terms are OFFSETS that add; the price is scaled by
  1 + their sum, floored at 0. Mirrors priceMultiplier in src/global.lua."""
  return max(0.0, 1 + sum(offsets))


def check_resale(rules, prices):
  """Warn, per rep band, when buying an item and selling it straight back
  makes money.

  Allowed on purpose -- items are finite, so a band that pays more than it
  charges is a reward, not an exploit -- but it should never be a surprise.
  Rounded the way the board rounds (buy up, sell down), so an item whose
  profit rounds away is not counted.
  """
  m = rules["margin"]
  for band in REP_ORDER:
    buy_mult = trade_multiplier(m["buy_pct"], rules["rep_buy"][band])
    sell_mult = trade_multiplier(m["sell_pct"], rules["rep_sell"][band])
    if sell_mult < buy_mult:
      continue
    gains = []
    for name, base in prices.items():
      # round(.., 6) is the board's float snap: 100 x 1.0499... is 105.
      gain = (math.floor(round(base * sell_mult, 6))
              - math.ceil(round(base * buy_mult, 6)))
      if gain > 0:
        gains.append((gain, name))
    if not gains:
      continue
    gain, name = max(gains)
    warn("rep band %r charges %.2f%% and pays %.2f%% -- buy-and-resell "
         "profits on %d of %d items (most: %s, +%dg per loop)"
         % (band, buy_mult * 100, sell_mult * 100, len(gains), len(prices),
            name, gain))


def rounded(rules, value):
  """Always round DOWN to the nearest step, never below the floor."""
  step = rules["anchor"].get("round_to", 1) or 1
  low = int(rules["anchor"].get("min_gold", 0))
  return max(int(math.floor(value / step) * step), low)


# ---- item pricing ---------------------------------------------------------

def ladder_for(itype, subtype):
  """Which tier ladder prices this row.

  A tier is only meaningful against a ladder: a tier-2 ore, a tier-2 sword and
  a tier-2 profession are not worth the same, so each family gets its own.

    goods/material, goods/junk, goods/treasure   goods_tier
    goods/miscellaneous                          misc_tier
    goods/quest                                  quest_item_tier
    consumable                                   consumable_tier
    skill                                        skill_tier
    weapon, armor (incl. their misc subtypes)    equip_tier

  `miscellaneous` is a SUBTYPE on several types, and on weapon/armor it means
  a gathering tool -- real equipment with a real tier ladder. Only
  goods/miscellaneous (a container, an oddment) takes misc_tier, which is why
  this dispatches on the pair rather than on the subtype alone.

  A lootbox takes no ladder here: lootbox_tier is the gold it DROPS, and its
  own price derives from that in lootbox_price().
  """
  if itype == "goods":
    if subtype == "miscellaneous":
      return "misc_tier"
    if subtype == "quest":
      return "quest_item_tier"
    return "goods_tier"
  if itype == "consumable":
    return "consumable_tier"
  if itype == "skill":
    return "skill_tier"
  return "equip_tier"


def hands_of(subsubtype):
  """subsubtype carries 1H/2H for weapons; anything else has no hand cost."""
  s = (subsubtype or "").strip().upper()
  if s.startswith("2H"):
    return 2
  if s.startswith("1H"):
    return 1
  return None


def lootbox_gold(rules, tier):
  """Gold a box of this tier DROPS -- the midpoint its spawngold rolls around."""
  return rules["anchor"]["gold_anchor"] * rules["lootbox_tier"][str(tier)]


def lootbox_price(rules, tier):
  """base_gold for the sealed box, set so SELLING it returns
  lootbox/sell_pct of the gold it would have dropped.

  A vendor pays (1 + margin/sell_pct) of base_gold like any other item, so
  the stored value has to be grossed up by that margin -- otherwise the two
  percentages would compound and the box would fetch far less than intended.
  """
  want = lootbox_gold(rules, tier) * rules["lootbox"]["sell_pct"]
  return rounded(rules, want / trade_multiplier(rules["margin"]["sell_pct"]))


def item_value(rules, itype, subtype, subsubtype, tier, hands):
  """Intrinsic worth at neutral reputation. None means the item has no price,
  which happens only when the row has no tier -- a tier is what a ladder
  prices against.
  """
  if tier is None:
    return None
  if subtype == "lootbox":
    return lootbox_price(rules, tier)

  ladder = ladder_for(itype, subtype)
  pct = rules[ladder].get(str(tier))
  if pct is None:
    warn("no %s entry for tier %s (%s/%s)" % (ladder, tier, itype, subtype))
    return None

  value = rules["anchor"]["gold_anchor"] * pct
  value *= rules["equip_type"].get(itype, 1.0)
  if hands:
    value *= rules["equip_hands"].get(str(hands), 1.0)
  # Treasure rides the same tier ladder as any other material, then takes a
  # premium: a gem has no crafting use, so selling it is what it is FOR.
  if (subsubtype or "").strip().lower() in TREASURE_SUBSUBTYPES:
    value *= rules["treasure"]["gem_mult"]
  # Junk is a trade good nobody wants: half of the real thing at its tier.
  # At low tiers that floors to 0, which is correct -- a rock is a rock.
  if subtype == "junk":
    value *= rules["junk"]["junk_mult"]
  return rounded(rules, value)


def sync_items(rules, check, force):
  """Price every item sheet; they share headers, so each goes through the
  same pass. One load and one save, because Excel locks the whole file."""
  wb = load_sheet(TILES)
  priced = stamped = 0
  rows = []
  for sheet in ITEM_SHEETS:
    p, s = price_sheet(wb[sheet], rules, check, force, rows)
    priced += p
    stamped += s
  if not check:
    save_sheet(wb, TILES)
  return priced, stamped, rows


def price_sheet(ws, rules, check, force, rows):
  ix = header_index(ws)
  for need in ("name", "type", "subtype", "subsubtype", "tier", "base_gold"):
    if need not in ix:
      raise SystemExit("tiles.xlsx %s: no %r column" % (ws.title, need))

  gold_col = ix["base_gold"] + 1
  mark_computed(ws, gold_col)
  gm_col = ix["gmnotes"] + 1 if "gmnotes" in ix else None
  if gm_col:
    mark_computed(ws, gm_col)

  priced = 0
  stamped = 0
  for row in range(2, ws.max_row + 1):
    name = ws.cell(row=row, column=ix["name"] + 1).value
    if not name:
      continue
    itype = (ws.cell(row=row, column=ix["type"] + 1).value or "").strip()
    subtype = (ws.cell(row=row, column=ix["subtype"] + 1).value or "").strip()
    subsub = ws.cell(row=row, column=ix["subsubtype"] + 1).value
    raw_tier = ws.cell(row=row, column=ix["tier"] + 1).value
    try:
      tier = int(raw_tier) if raw_tier not in (None, "") else None
    except (TypeError, ValueError):
      warn("%s: tier %r is not a number" % (name, raw_tier))
      tier = None

    # gmnotes is stamped for the families whose mechanics are fully derived:
    # a material's kind+tier (what a recipe matches on) and a weapon's damage
    # line (what a combat script reads). Everything else stays authored.
    if gm_col:
      want_gm = None
      if itype == "goods" and subtype == "material":
        if subsub and tier is not None:
          want_gm = "%s%d" % (str(subsub).strip(), tier)
        else:
          warn("%s: goods/material needs subsubtype and tier for gmnotes"
               % name)
      elif itype == "weapon":
        want_gm = weapon_gmnotes(subtype, tier, hands_of(subsub))
        if want_gm is None:
          warn("%s: weapon has no derivable damage for gmnotes (tier %r, "
               "subtype %r)" % (name, tier, subtype))

      if want_gm is not None:
        gm = ws.cell(row=row, column=gm_col)
        if gm.value != want_gm:
          if not check:
            gm.value = want_gm
          stamped += 1

    cell = ws.cell(row=row, column=gold_col)
    want = item_value(rules, itype, subtype, subsub, tier, hands_of(subsub))

    # A non-numeric value here is a typo (a URL landed in this column once),
    # not an override -- say so loudly rather than preserving it.
    cur = cell.value
    if cur not in (None, "") and not isinstance(cur, (int, float)):
      warn("%s: base_gold held non-numeric %r -- replaced"
           % (name, str(cur)[:40]))
      cur = None

    if cur not in (None, "") and want is not None and int(cur) != want:
      if force:
        overwritten_out.append("%s: %dg -> %dg" % (name, int(cur), want))
      else:
        overrides_out.append("%s: %dg authored, rules say %dg"
                             % (name, int(cur), want))
        rows.append((name, itype, subtype, tier, int(cur), True))
        continue

    if not check:
      cell.value = want
    if want is not None:
      priced += 1
    rows.append((name, itype, subtype, tier, want, False))
  return priced, stamped


# ---- origins --------------------------------------------------------------

def sync_origins(rules, check, force):
  """spawngold comes from the origin's gold_tier. Drift against the prose
  "Gain N gold." is reported, since this tool is what causes it."""
  wb = load(ORIGINS, data_only=False)
  rows = read_rows(ORIGINS, ORIGINS_SHEET, "origin_id", wb)
  if not rows:
    return 0, []
  if "gold_tier" not in rows[0]:
    warn("origins.xlsx has no gold_tier column -- add one (1-5) and re-run; "
         "spawngold left untouched")
    return 0, []

  anchor = rules["anchor"]["gold_anchor"]
  out = []
  changed = 0
  for r in rows:
    raw = (r.get("gold_tier") or "").strip()
    pct = rules["origin_tier"].get(raw) if raw else None
    if raw and pct is None:
      warn("origins.xlsx %s: unknown gold_tier %r" % (r["origin_id"], raw))
      out.append((r["origin_id"], r.get("name", ""), None))
      continue
    if pct is None:
      if (r.get("spawngold") or "").strip():
        changed += 1
      r["spawngold"] = ""
      out.append((r["origin_id"], r.get("name", ""), None))
      continue

    value = rounded(rules, anchor * pct)
    if (r.get("spawngold") or "").strip() != str(value):
      changed += 1
    r["spawngold"] = str(value)

    # The effect text names a number the player reads on the card.
    said = re.search(r"Gain (\d+) gold", r.get("effect") or "")
    if said and int(said.group(1)) != value:
      warn("origins.xlsx %s (%s): text says %s gold, rules give %d"
           % (r["origin_id"], r.get("name", ""), said.group(1), value))
    out.append((r["origin_id"], r.get("name", ""), value))

  if not check and changed:
    write_back(wb, ORIGINS_SHEET, rows, ("spawngold",))
    save(wb, ORIGINS)
  return changed, out


# ---- lootboxes ------------------------------------------------------------

def item_prices():
  """Every priced item, name -> base_gold, read back from the sheet so the
  splice reflects hand overrides as well as computed values."""
  wb = load_sheet(TILES)
  prices = {}
  for sheet in ITEM_SHEETS:
    ws = wb[sheet]
    ix = header_index(ws)
    for row in ws.iter_rows(min_row=2, values_only=True):
      name = row[ix["name"]]
      gold = row[ix["base_gold"]]
      if name and isinstance(gold, (int, float)):
        prices[str(name)] = int(gold)
  return prices


def item_tiers():
  """Tier lives on the item sheets of tiles.xlsx, keyed by name."""
  wb = load_sheet(TILES)
  tiers = {}
  for sheet in ITEM_SHEETS:
    ws = wb[sheet]
    ix = header_index(ws)
    for row in ws.iter_rows(min_row=2, values_only=True):
      name = row[ix["name"]]
      if name:
        tiers[name] = row[ix["tier"]]
  return tiers


def sync_lootboxes(rules, check, force):
  """Write spawngold for the boxes that carry coin.

  Not every "lootbox" is a container of money. A clam is an organism whose
  worth is the pearl inside it, and a crab trap yields livestock -- both are
  lootbox-shaped (they spawn items on opening) but neither holds coin. The
  gold_adjust column marks which rows this tool may touch:

    1  this box pays coin -- spawngold is COMPUTED here
    0  this box pays items only -- spawngold is left exactly as authored

  A 0 row is not "zero gold" by decree; it is hand-owned, so a designer can
  still type a one-off amount without the sync overwriting it.
  """
  wb = load_sheet(TILES)
  ws = wb[LOOTBOX_SHEET]
  ix = header_index(ws)
  if "spawngold" not in ix:
    warn("lootbox_table has no spawngold column")
    return 0, []
  if "gold_adjust" not in ix:
    warn("lootbox_table has no gold_adjust column -- add one (0/1) and "
         "re-run; spawngold left untouched")
    return 0, []

  tiers = item_tiers()
  gcol = ix["spawngold"] + 1
  mark_computed(ws, gcol)

  var = rules["lootbox"]["variance"]
  out = []
  changed = 0
  for row in range(2, ws.max_row + 1):
    name = ws.cell(row=row, column=ix["name"] + 1).value
    if not name:
      continue

    flag = ws.cell(row=row, column=ix["gold_adjust"] + 1).value
    if flag in (None, ""):
      warn("lootbox_table: %s has no gold_adjust -- skipped (set 0 or 1)"
           % name)
      out.append((name, None, None, "skipped"))
      continue
    try:
      adjust = int(flag)
    except (TypeError, ValueError):
      warn("lootbox_table: %s gold_adjust is %r, want 0 or 1" % (name, flag))
      out.append((name, None, None, "skipped"))
      continue
    if not adjust:
      cur = ws.cell(row=row, column=gcol).value
      out.append((name, None, None, "hand-owned: %s" % (cur or "none")))
      continue

    raw = tiers.get(name)
    try:
      tier = int(raw) if raw not in (None, "") else None
    except (TypeError, ValueError):
      tier = None
    if tier is None:
      warn("lootbox_table: %s has no tier on an item sheet" % name)
      out.append((name, None, None, None))
      continue

    mid = lootbox_gold(rules, tier)
    lo = rounded(rules, mid * (1 - var))
    hi = rounded(rules, mid * (1 + var))
    text = "%d-%d" % (lo, hi)
    if (ws.cell(row=row, column=gcol).value or "") != text:
      changed += 1
    if not check:
      ws.cell(row=row, column=gcol).value = text
    out.append((name, tier, rounded(rules, mid), text))

  if not check:
    save_sheet(wb, TILES)
  return changed, out


# ---- Lua output -----------------------------------------------------------

def write_lua(rules, prices, splice=True):
  """Margins, rep bands and live prices go to Global.

  Vendors are interactive and Global owns every click handler, so the table a
  vendor reads lives there -- the same argument that keeps EVENT_DATA in
  Global rather than on the Card Importer bag.

  ECONOMY.prices duplicates base_gold from ITEM_DATA deliberately. ITEM_DATA
  is inlined into the importer bag, so a price change there only reaches
  items imported AFTERWARDS. Reading prices from Global instead means a
  re-sync plus CTRL+SHIFT+E reprices every item already on the table.
  """
  anchor = rules["anchor"]["gold_anchor"]
  lines = ["-- GENERATED by tools/sync/sync_gold.py -- do not edit by hand",
           "ECONOMY = {",
           "  anchor = %d," % int(anchor),
           "  party_size = %d," % int(rules["anchor"].get("party_size", 1)),
           "  gem_mult = %g," % rules["treasure"]["gem_mult"],
           "  round_to = %d," % int(rules["anchor"].get("round_to", 1) or 1),
           "  min_gold = %d," % int(rules["anchor"].get("min_gold", 0)),
           "  sell_pct = %g," % rules["margin"]["sell_pct"],
           "  buy_pct = %g," % rules["margin"]["buy_pct"],
           "  craft_mult = %g," % rules["craft"]["value_mult"],
           "  craft_floor = %d," % int(rules["craft"]["value_floor"]),
           "  rep_order = {%s},"
           % ", ".join('"%s"' % n for n in REP_ORDER),
           "  rep = {"]
  for n in REP_ORDER:
    lines.append('    {name = "%s", min = %d, buy = %g, sell = %g},'
                 % (n, int(rules["rep"][n]), rules["rep_buy"][n],
                    rules["rep_sell"][n]))
  lines.append("  },")
  for g in TIER_LADDERS:
    vals = ", ".join("[%d] = %d" % (t, rounded(rules, anchor * rules[g][str(t)]))
                     for t in range(4))
    lines.append("  %s = {%s}," % (g.replace("_tier", "_gold"), vals))

  # Live prices, keyed by the name findInBag already matches on.
  lines.append("  prices = {")
  for name in sorted(prices):
    lines.append('    [%s] = %d,' % (lua_str(name), prices[name]))
  lines.append("  },")
  lines.append("}")
  data_lua = "\n".join(lines)

  os.makedirs(os.path.dirname(OUT), exist_ok=True)
  with open(OUT, "w", encoding="utf-8", newline="\n") as f:
    f.write(data_lua + "\n")

  if not splice:
    return 0

  #[[ Splice into the Global that sync_events_data.py rebuilt from
  # src/global.lua. Running against a stale one would produce a file with the
  # data but none of the functions that use it -- same guard as
  # sync_origins_data.py.
  if not os.path.exists(GLOBAL):
    warn("%s missing -- ECONOMY not spliced; run tools/sync/sync_events_data.py "
         "first" % os.path.relpath(GLOBAL, WORKSPACE))
    return 0
  with open(GLOBAL, encoding="utf-8") as f:
    script = f.read()
  if "function renderCard" not in script:
    warn("%s was not built from src/global.lua -- ECONOMY not spliced"
         % os.path.relpath(GLOBAL, WORKSPACE))
    return 0

  block = BEGIN + "\n" + data_lua + "\n" + END
  if BEGIN in script and END in script:
    script = script.split(BEGIN)[0] + block + script.split(END, 1)[1]
  else:
    script = script.rstrip() + "\n\n" + block + "\n"
  with open(GLOBAL, "w", encoding="utf-8", newline="\n") as f:
    f.write(script)
  return len(script)


def lua_str(s):
  s = (s or "").replace("\\", "\\\\").replace('"', '\\"')
  s = s.replace("\r", "").replace("\n", "\\n")
  return '"' + s + '"'


# ---- main -----------------------------------------------------------------

def main():
  check = "--check" in sys.argv
  force = "--force" in sys.argv
  rules = read_rules()
  validate(rules)

  #[[ --splice-only is the SECOND pass in master_sync.cmd. The first pass
  # writes the sheets, then sync_events_data.py rebuilds Global.lua from
  # src/global.lua (wiping any earlier splice), so ECONOMY has to go in after
  # that. Re-reading the sheets here costs nothing and keeps one source.
  if "--splice-only" in sys.argv:
    prices = item_prices()
    check_resale(rules, prices)
    size = write_lua(rules, prices)
    if size:
      print("sync_gold: spliced ECONOMY into %s (%d prices, %d bytes)"
            % (os.path.relpath(GLOBAL, WORKSPACE), len(prices), size))
    for w in warnings_out:
      print("  warning: " + w)
    return 0

  priced, stamped, _ = sync_items(rules, check, force)
  o_changed, _ = sync_origins(rules, check, force)
  l_changed, _ = sync_lootboxes(rules, check, force)
  if not check:
    # Global.lua does not exist yet on a clean run; the --splice-only pass
    # puts ECONOMY in once sync_events_data.py has built it.
    write_lua(rules, item_prices(), splice=False)
  check_resale(rules, item_prices())

  verb = "would price" if check else "priced"
  print("sync_gold: %s %d items, %d origins, %d lootboxes, "
        "%d gmnotes (anchor %d)"
        % (verb, priced, o_changed, l_changed, stamped,
           int(rules["anchor"]["gold_anchor"])))
  for o in overwritten_out:
    print("  overwritten (--force): " + o)
  for o in overrides_out:
    print("  OVERRIDE kept: " + o)
  for w in warnings_out:
    print("  warning: " + w)
  return 0


if __name__ == "__main__":
  sys.exit(main())
