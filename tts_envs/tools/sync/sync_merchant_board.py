#!/usr/bin/env python3
"""Inline shops.xlsx into the Merchant Board script and stage it.

  data/shops.xlsx   ->  SHOPS + STOCK  ->  src/merchant_board.lua
  data/tokens.xlsx                     ->  .tts/objects/<nickname>.ecde97.lua
                    ->  FACTIONS       ->  .tts/objects/Global.lua

A shop's tier and faction belong to the merchant TOKEN, not the shop: they
come from the tokens.xlsx row whose role is `shop` and whose gmnotes is the
shop_id -- the same GMNotes the board reads off the meeple in its slot.

SHOPS and STOCK live on the BOARD, not in Global, because the board is their
only consumer -- the same reason ITEM_DATA sits on the importer bag. FACTIONS
goes the other way, into Global: reputation is priced there, and a faction's
starting standing is part of that.

Every item name and faction id is checked here. A typo or a renamed item
fails the build rather than appearing in game as a ware nothing can price --
the failure mode that cost us 'Waterlogged Pack' once already.

Run AFTER tools/sync/sync_events_data.py, which rebuilds Global.lua from
src/global.lua and would wipe an earlier FACTIONS splice.

Run: python tools/sync/sync_merchant_board.py
"""
import os, re, sys

sys.dont_write_bytecode = True
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(1, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))  # tools/: datasets, spawn_cells, lua_inline
from sync_change_converter import check_lua52
from spawn_cells import staged_path, strip_unique, weighted_list
from datasets import (ROOT, SHOPS, TILES, TOKENS, TOKEN_SHEET,
                      FACTIONS_SHEET, SHOP_ROLE, item_names, read_rows)
from sync_gold import REP_ORDER, read_rules, trade_multiplier

SRC = os.path.join(ROOT, "src", "merchant_board.lua")
OUT = os.path.join(ROOT, "generated", "merchant_data.lua")
WORKSPACE = os.path.dirname(ROOT)
GLOBAL = os.path.join(WORKSPACE, ".tts", "objects", "Global.lua")
GUID = "ecde97"

BEGIN = "--[[ BEGIN GENERATED SHOP DATA -- do not edit ]]"
END = "--[[ END GENERATED SHOP DATA ]]"
FACTION_BEGIN = "--[[ BEGIN GENERATED FACTION DATA -- do not edit ]]"
FACTION_END = "--[[ END GENERATED FACTION DATA ]]"

TIERS = (0, 1, 2, 3)    # settlement, town, city, capital
PLACEHOLDER = "TBD:"    # stock for an item that is not in tiles.xlsx yet

warnings_out = []


def warn(msg):
  warnings_out.append(msg)


def load_sheet(path):
  import warnings
  warnings.filterwarnings("ignore", message=".*Data Validation extension.*")
  from openpyxl import load_workbook
  return load_workbook(path)


def rows_of(wb, sheet):
  """Dicts keyed by header, skipping blank rows."""
  if sheet not in wb.sheetnames:
    raise SystemExit("shops.xlsx: no %r sheet" % sheet)
  ws = wb[sheet]
  header = [c.value for c in ws[1]]
  out = []
  for row in ws.iter_rows(min_row=2, values_only=True):
    if all(v in (None, "") for v in row):
      continue
    out.append({h: v for h, v in zip(header, row) if h})
  return out


def text(v):
  return "" if v is None else str(v).strip()


def adjustment(s, col, where):
  """A shop's own price OFFSET (-0.1 = 10% less). Blank means 0."""
  raw = s.get(col)
  if raw in (None, ""):
    return 0.0
  try:
    return float(raw)
  except (TypeError, ValueError):
    raise SystemExit("%s: %s %r is not a number" % (where, col, raw))


def check_spread(sid, shop, rules):
  """Warn when a shop pays more for an item than it charges for it.

  buy and sell are adjusted separately, so nothing forces the gap between
  them to stay open the way the old single markup did. Party bonuses are not
  counted -- they do not exist yet.
  """
  m = rules["margin"]
  for band in REP_ORDER:
    buy = trade_multiplier(m["buy_pct"], rules["rep_buy"][band],
                           shop["adjustment_buy"])
    sell = trade_multiplier(m["sell_pct"], rules["rep_sell"][band],
                            shop["adjustment_sell"])
    if sell >= buy:
      warn("%s (%s) at %s: sells to the player at %.0f%% and buys back at "
           "%.0f%% -- buy-and-resell is profitable"
           % (sid, shop["role"], band, buy * 100, sell * 100))
      return


def lua_str(s):
  s = text(s).replace("\\", "\\\\").replace('"', '\\"')
  s = s.replace("\r", "").replace("\n", "\\n")
  return '"' + s + '"'


def parse_qty(cell, where):
  """A count or a range: 2 -> (2, 2), "1-3" -> (1, 3), blank -> (1, 1).

  A range starting at 0 is how a row says "sometimes absent". Excel turns a
  bare 2 into a number and a 1-3 into text, so both arrive here.
  """
  if cell is None or text(cell) == "":
    return 1, 1
  if isinstance(cell, (int, float)) and not isinstance(cell, bool):
    n = int(cell)
    lo = hi = n
  else:
    m = re.fullmatch(r"\s*(\d+)\s*(?:-\s*(\d+)\s*)?", text(cell))
    if not m:
      raise SystemExit("%s: qty %r -- want a number (2) or a range (1-3)"
                       % (where, cell))
    lo = int(m.group(1))
    hi = int(m.group(2)) if m.group(2) else lo
  if lo > hi:
    raise SystemExit("%s: qty %r runs backwards" % (where, cell))
  if hi == 0:
    raise SystemExit("%s: qty %r can never produce anything" % (where, cell))
  return lo, hi


def splice(path, begin, end, block, guard=None):
  with open(path, encoding="utf-8") as f:
    script = f.read()
  if guard and guard not in script:
    raise SystemExit("%s was not built from src/global.lua -- run "
                     "tools/sync/sync_events_data.py first"
                     % os.path.relpath(path, WORKSPACE))
  full = begin + "\n" + block + "\n" + end
  if begin in script and end in script:
    script = script.split(begin)[0] + full + script.split(end, 1)[1]
  else:
    script = script.rstrip() + "\n\n" + full + "\n"
  with open(path, "w", encoding="utf-8", newline="\n") as f:
    f.write(script)
  return len(script)


def main():
  for path in (SHOPS, TILES, TOKENS, SRC, GLOBAL):
    if not os.path.exists(path):
      raise SystemExit("missing %s" % os.path.relpath(path, WORKSPACE))

  wb = load_sheet(SHOPS)
  factions_rows = rows_of(load_sheet(TOKENS), FACTIONS_SHEET)
  shops = rows_of(wb, "shops")
  stock = rows_of(wb, "stock")
  known = item_names()

  # ---- factions ----------------------------------------------------------
  factions = {}
  for f in factions_rows:
    fid = text(f.get("faction_id"))
    if not fid:
      raise SystemExit("tokens.xlsx factions: a row has no faction_id")
    if fid in factions:
      raise SystemExit("tokens.xlsx factions: duplicate faction_id %r" % fid)
    try:
      rep = int(f.get("default_rep") or 0)
    except (TypeError, ValueError):
      raise SystemExit("tokens.xlsx factions: %s default_rep %r is not a "
                       "number" % (fid, f.get("default_rep")))
    factions[fid] = {"name": text(f.get("name")) or fid, "rep": rep}
  if "none" not in factions:
    raise SystemExit("tokens.xlsx factions: need a 'none' row for the "
                     "Woodland's unaligned traders")

  # ---- shop tokens -------------------------------------------------------
  # shop_id -> its merchant token. sync_items_data.py validates the token
  # sheet in full; only what the board needs is checked again here.
  token_of = {}
  for t in read_rows(TOKENS, TOKEN_SHEET, "name"):
    if t.get("role", "").lower() != SHOP_ROLE:
      continue
    sid = t.get("gmnotes", "")
    if sid in token_of:
      raise SystemExit("%s: shop_id %r is already the gmnotes of %s"
                       % (t["_where"], sid, token_of[sid]["_where"]))
    token_of[sid] = t

  # ---- shops -------------------------------------------------------------
  by_id = {}
  for s in shops:
    sid = text(s.get("shop_id"))
    role = text(s.get("shop_role_display"))
    if not sid:
      raise SystemExit("shops.xlsx shops: a row has no shop_id")
    # A prefilled id with no name yet is not a shop.
    if not role:
      continue
    if sid in by_id:
      raise SystemExit("shops.xlsx shops: duplicate shop_id %r" % sid)
    token = token_of.get(sid)
    if token is None:
      raise SystemExit("shops.xlsx shops: %s (%s) has no merchant token -- "
                       "add a tokens.xlsx row with role 'shop' and gmnotes "
                       "%r" % (sid, role, sid))
    where = "%s (%s)" % (token["_where"], token["name"])
    faction = token.get("faction") or "none"
    if faction not in factions:
      raise SystemExit("%s: faction %r is not on the factions sheet"
                       % (where, faction))
    try:
      tier = int(token.get("tier"))
    except (TypeError, ValueError):
      raise SystemExit("%s: a shop token needs a tier (0-3)" % where)
    if tier not in TIERS:
      raise SystemExit("%s: tier %r -- want 0-3" % (where, tier))
    if "markup" in s:
      raise SystemExit("shops.xlsx shops: the markup column was split into "
                       "adjustment_buy and adjustment_sell")
    by_id[sid] = {"role": role, "tier": tier, "faction": faction,
                  "adjustment_buy": adjustment(s, "adjustment_buy", where),
                  "adjustment_sell": adjustment(s, "adjustment_sell", where)}
  for sid in sorted(set(token_of) - set(by_id)):
    raise SystemExit("%s: gmnotes %r is not a shop_id in shops.xlsx"
                     % (token_of[sid]["_where"], sid))
  rules = read_rules()
  for sid in by_id:
    check_spread(sid, by_id[sid], rules)

  # ---- stock -------------------------------------------------------------
  rows_by_shop = {}
  placeholders = {}
  for n, r in enumerate(stock, 2):
    where = "shops.xlsx stock row %d" % n
    sid = text(r.get("shop_id"))
    if sid not in by_id:
      raise SystemExit("%s: unknown shop_id %r" % (where, sid))
    # Same grammar as lootboxes: "Pickaxe|Woodaxe| " picks one per Restock.
    entries = weighted_list(text(r.get("item")))
    if not entries:
      continue
    #[[ An unknown name is fatal, not a warning. The board would spawn
    # nothing for it and the slot would silently come up empty, which reads
    # as "the merchant is out of stock" rather than as a typo.
    #
    # The one exception is a deliberate PLACEHOLDER, "TBD:<description>", for
    # an item that is designed but not authored yet. It is reported and
    # emitted as a blank pick, so in game it simply never appears.
    for i, entry in enumerate(entries):
      name = strip_unique(entry)
      if name.startswith(PLACEHOLDER):
        placeholders.setdefault(name, set()).add(sid)
        entries[i] = ""
      elif entry and name not in known:
        raise SystemExit("%s: %s wants %r, which is not in tiles.xlsx -- "
                         "write %s%s if it is not authored yet"
                         % (where, sid, name, PLACEHOLDER, name))
    if not any(entries):
      continue
    lo, hi = parse_qty(r.get("qty"), where)
    rows_by_shop.setdefault(sid, []).append(
        {"items": entries, "lo": lo, "hi": hi})

  if placeholders:
    shops_with = set().union(*placeholders.values())
    warn("%d placeholder item%s (%s...) across %d shops -- they spawn nothing "
         "until added to tiles.xlsx; see plans/auto_TODO.md"
         % (len(placeholders), "" if len(placeholders) == 1 else "s",
            PLACEHOLDER, len(shops_with)))

  for sid in by_id:
    if not rows_by_shop.get(sid):
      warn("%s (%s) has no stock -- Restock will find nothing"
           % (sid, by_id[sid]["role"]))

  # ---- emit: the board ---------------------------------------------------
  def idkey(sid):
    m = re.search(r"\d+", sid)
    return (int(m.group()) if m else 0, sid)

  L = ["-- GENERATED by tools/sync/sync_merchant_board.py -- do not edit by hand",
       "SHOPS = {"]
  for sid in sorted(by_id, key=idkey):
    s = by_id[sid]
    L.append("  [%s] = {role = %s, tier = %d, faction = %s, "
             "adjustment_buy = %g, adjustment_sell = %g},"
             % (lua_str(sid), lua_str(s["role"]), s["tier"],
                lua_str(s["faction"]), s["adjustment_buy"],
                s["adjustment_sell"]))
  L.append("}")
  L.append("")
  L.append("-- What each merchant MAY carry, in slot order. One of items is")
  L.append("-- picked per Restock (\"\" = nothing); qty rolls between lo and hi.")
  L.append("STOCK = {")
  for sid in sorted(rows_by_shop, key=idkey):
    L.append("  [%s] = {" % lua_str(sid))
    for r in rows_by_shop[sid]:
      L.append("    {items = {%s}, lo = %d, hi = %d},"
               % (", ".join(lua_str(e) for e in r["items"]), r["lo"], r["hi"]))
    L.append("  },")
  L.append("}")
  data_lua = "\n".join(L)

  os.makedirs(os.path.dirname(OUT), exist_ok=True)
  with open(OUT, "w", encoding="utf-8", newline="\n") as f:
    f.write(data_lua + "\n")

  with open(SRC, encoding="utf-8") as f:
    script = f.read()
  check_lua52(script, os.path.relpath(SRC, ROOT))
  block = BEGIN + "\n" + data_lua + "\n" + END
  if BEGIN in script and END in script:
    script = script.split(BEGIN)[0] + block + script.split(END, 1)[1]
  else:
    script = block + "\n\n" + script
  dest = staged_path(GUID)
  with open(dest, "w", encoding="utf-8", newline="\n") as f:
    f.write(script)

  # ---- emit: Global ------------------------------------------------------
  F = ["-- GENERATED by tools/sync/sync_merchant_board.py -- do not edit by hand",
       "-- Every faction a merchant can belong to, and where the party starts",
       "-- with them before anything has happened.",
       "FACTIONS = {"]
  for fid in sorted(factions):
    F.append("  [%s] = {name = %s, default_rep = %d},"
             % (lua_str(fid), lua_str(factions[fid]["name"]),
                factions[fid]["rep"]))
  F.append("}")
  gsize = splice(GLOBAL, FACTION_BEGIN, FACTION_END, "\n".join(F),
                 guard="function renderCard")

  print("sync_merchant_board: %d shops, %d stock rows -> %s (%d bytes)"
        % (len(by_id), sum(len(v) for v in rows_by_shop.values()),
           os.path.relpath(dest, WORKSPACE), len(script)))
  print("sync_merchant_board: %d factions -> %s (%d bytes)"
        % (len(factions), os.path.relpath(GLOBAL, WORKSPACE), gsize))
  for w in warnings_out:
    print("  warning: " + w)
  return 0


if __name__ == "__main__":
  sys.exit(main())
