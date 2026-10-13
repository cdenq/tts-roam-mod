"""Where the authored data lives, and how a sheet becomes a list of rows.

Every sync reads data/ through these names, so moving a file or adding a sheet
is one edit here instead of a hunt through every tool.

  data/tiles.xlsx           items, skills, statuses  one row per tile/token
                            lootbox_table          what each lootbox drops
  data/tokens.xlsx          tokens                   one row per meeple token
                            factions                 every faction, and its starting rep
                            loot_table               what a token drops when killed
                            stats_table              a token's combat numbers
                            landmarks                places on the map: caves, shrines, camps
  data/scenarios.xlsx       scenario_intro           one row per event card
                            scenario_options         the buttons on each card
                            scenario_outcomes        what each button can lead to
  data/origins.xlsx         origins                  one row per origin card
  data/gold_variables.csv
  data/shops.xlsx
  data/asset_variables.csv  art knobs for tools/build_map_tiles.py + build_gems.py + build_landmarks.py
  data/maptile.xlsx         <deck> x5, backs       image URLs for the Maptile Manager
                            biome_descriptions     tile description per biome
                            maptile_table          what Render spawns, per biome (GMNotes)
                            density, weights       explorations per density; draw odds
"""
import os, re, sys

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DATA = os.path.join(ROOT, "data")

TILES = os.path.join(DATA, "tiles.xlsx")
TOKENS = os.path.join(DATA, "tokens.xlsx")
SCENARIOS = os.path.join(DATA, "scenarios.xlsx")
ORIGINS = os.path.join(DATA, "origins.xlsx")
GOLD = os.path.join(DATA, "gold_variables.csv")
SHOPS = os.path.join(DATA, "shops.xlsx")
ASSET_VARS = os.path.join(DATA, "asset_variables.csv")
MAPTILES = os.path.join(DATA, "maptile.xlsx")

#[[ Sheets that share one set of headers and are read as ONE table, in this
# order. Adding a tile family means adding a sheet with the same headers and
# naming it here -- nothing else changes.
ITEM_SHEETS = ("items", "skills", "statuses")
LOOTBOX_SHEET = "lootbox_table"

#[[ Tokens are meeples: people, not things. They import through the same bag
# as tiles, so a token name must not collide with any tile name. A shop
# token's gmnotes is the shop_id the Merchant Board looks up in shops.xlsx.
TOKEN_SHEET = "tokens"
#[[ Landmarks are tokens too (same bag, same name space, same importer), with
# fewer columns: no role/affiliation/exp/loot/stats. They import as role
# "landmark" so ITEM_DATA and the description builder need nothing new.
LANDMARK_SHEET = "landmarks"
LANDMARK_ROLE = "landmark"
FACTIONS_SHEET = "factions"
LOOT_TABLE_SHEET = "loot_table"
STATS_TABLE_SHEET = "stats_table"
SHOP_ROLE = "shop"
COMBAT_ROLE = "combat"
ORIGINS_SHEET = "origins"


#[[ A scenario is an intro card, its options, and the outcomes those options
# roll. Region is a column on the intro, not a sheet, so every region shares
# these three. An outcome is keyed by scenario_id + outcome_key; outcome_id()
# joins them into the node id the game stores ("EV1_pay_through").
SCENARIO_INTRO = "scenario_intro"
SCENARIO_OPTIONS = "scenario_options"
SCENARIO_OUTCOMES = "scenario_outcomes"


def outcome_id(scenario_id, outcome_key):
  return "%s_%s" % (scenario_id, outcome_key)


def load(path, data_only=True):
  import warnings
  #[[ openpyxl cannot represent Excel's data-validation EXTENSION and says so
  # on every read. The warning is pure noise that buries real errors.
  warnings.filterwarnings("ignore", message=".*Data Validation extension.*")
  try:
    from openpyxl import load_workbook
  except ImportError:
    sys.exit("openpyxl is required to read %s -- pip install openpyxl"
             % os.path.basename(path))
  if not os.path.exists(path):
    sys.exit("missing %s" % path)
  return load_workbook(path, data_only=data_only)


def cell_text(v):
  """Stringify a cell the way a CSV would have held it."""
  if v is None:
    return ""
  # Excel stores a bare integer as a float; 2.0 must not become "2.0"
  if isinstance(v, float) and v.is_integer():
    return str(int(v))
  return str(v).strip()


def read_rows(path, sheets, key, wb=None):
  """Rows of one or more same-shaped sheets as dicts, skipping a blank key.

  Every value is stringified, so callers see what a CSV would have given them.
  Each row carries `_where` ("tiles.xlsx skills row 7") for error messages and
  `_row` / `_sheet` for anything that writes back.
  """
  if isinstance(sheets, str):
    sheets = (sheets,)
  wb = wb or load(path)
  fname = os.path.basename(path)
  out = []
  for sheet in sheets:
    if sheet not in wb.sheetnames:
      sys.exit("%s has no sheet named %r (found: %s)"
               % (fname, sheet, ", ".join(wb.sheetnames)))
    rows = wb[sheet].iter_rows(values_only=True)
    try:
      header = [cell_text(h) for h in next(rows)]
    except StopIteration:
      continue
    for n, r in enumerate(rows, start=2):   # start=2: row 1 is the header
      rec = {h: cell_text(r[i] if i < len(r) else None)
             for i, h in enumerate(header) if h}
      if rec.get(key):
        rec["_row"] = n
        rec["_sheet"] = sheet
        rec["_where"] = "%s %s row %d" % (fname, sheet, n)
        out.append(rec)
  return out


def template_cell(row, where):
  """The `template` column: GMNotes of the bag object this row is cloned from.

  Any value works. An importer treats an object as a template when its GMNotes
  matches a template its data names, so each sync also checks that no
  imported object's GMNotes collides with one (see template_collisions).
  """
  t = row.get("template", "")
  if not t:
    sys.exit("%s: no template -- name the GMNotes of the object to clone"
             % where)
  # Tabs and newlines would break the card index wire format (cardIndexText).
  if re.search(r"[\t\r\n]", t):
    sys.exit("%s: template %r contains a tab or newline" % (where, t))
  return t


def template_collisions(templates, gmnotes, what):
  """Fail when an imported object would carry a template's GMNotes.

  The importer would mistake it for the template: "Take out all" would skip
  it, and an import could clone it instead of the real one.
  """
  for t in sorted(set(templates) & set(gmnotes)):
    sys.exit("template %r is also the GMNotes of %s %r -- rename the "
             "template" % (t, what, t))


def header_index(ws):
  """Header name -> 1-based column number."""
  return {str(c.value).strip(): c.column for c in ws[1] if c.value}


def write_back(wb, sheet, rows, columns):
  """Write `columns` of rows from read_rows() back to their own cells.

  A whole-number string goes back as a number, so Excel does not flag it as
  "number stored as text".
  """
  ws = wb[sheet]
  ix = header_index(ws)
  for r in rows:
    for col in columns:
      v = r.get(col, "")
      ws.cell(row=r["_row"], column=ix[col]).value = (
        int(v) if re.fullmatch(r"0|-?[1-9]\d*", v) else (v or None))


def save(wb, path):
  """Excel holds an exclusive lock on an open workbook, and openpyxl's
  PermissionError says nothing useful about why."""
  try:
    wb.save(path)
  except PermissionError:
    sys.exit("cannot write %s -- it is open in Excel. Close it and re-run."
             % os.path.basename(path))


def scenario_node_ids():
  """Every node id the game can see: intros and their outcomes."""
  wb = load(SCENARIOS)
  ids = {r["scenario_id"] for r in read_rows(SCENARIOS, SCENARIO_INTRO,
                                             "scenario_id", wb)}
  ids |= {outcome_id(r["scenario_id"], r["outcome_key"])
          for r in read_rows(SCENARIOS, SCENARIO_OUTCOMES, "outcome_key", wb)}
  return ids


def item_names():
  """Every tile/token name across ITEM_SHEETS, for spawn cross-checks."""
  return {r["name"] for r in read_rows(TILES, ITEM_SHEETS, "name")}


def token_names():
  """Every meeple and landmark name in tokens.xlsx, for spawnmeeple cross-checks."""
  wb = load(TOKENS)
  sheets = [TOKEN_SHEET] + ([LANDMARK_SHEET] if LANDMARK_SHEET in wb.sheetnames else [])
  return {r["name"] for s in sheets for r in read_rows(TOKENS, s, "name", wb)}
